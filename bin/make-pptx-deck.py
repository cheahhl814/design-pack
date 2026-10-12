#!/usr/bin/env python3
"""make-pptx-deck.py: build a .pptx deck whose figures are svg-library components
embedded as native vector SVG (PowerPoint 2016+/365, Mac and Online), with a
high-res PNG fallback for clients that ignore SVG (Google Slides, Keynote, old
Office). All text stays native: titles in real title placeholders, bullets in
body placeholders, captions in caption placeholders, notes in the notes page.

    python3 bin/make-pptx-deck.py spec.json -o deck.pptx [--theme light|dark] [--base corp.pptx]

  spec     JSON (or YAML if PyYAML is installed), see templates/deck/README.md:
             {"title": "...", "theme": "light",
              "slides": [
                {"layout": "title", "title": "...", "subtitle": "..."},
                {"layout": "section", "title": "...", "subtitle": "..."},
                {"layout": "figure", "title": "...", "figure": "food-web",
                 "caption": "...", "alt": "...", "notes": "..."},
                {"layout": "two-figures", "title": "...",
                 "figures": [{"svg": "ocean-zones", "caption": "..."}, "water-cycle"]},
                {"layout": "bullets", "title": "...", "bullets": ["a", ["sub", "sub"], "b"]},
                {"layout": "agenda", "items": ["...", "..."]},
                {"layout": "quote", "quote": "...", "attribution": "Name, role"},
                {"layout": "table", "title": "...", "header": ["A", "B"], "rows": [["1", "2"]]},
                {"layout": "kpi", "title": "...", "value": "92%", "label": "...", "delta": "+12 pts",
                 "stats": [{"value": "48 h", "label": "..."}]},
                {"layout": "image-text", "title": "...", "figure": "food-web", "side": "right",
                 "bullets": ["..."]},
                {"layout": "closing", "contact": ["Questions?", "team@example.org"]}],
              "footer": "Deck title", "date": "auto", "slide_numbers": true,
              "credits": ["deck-level credit lines, added to the last slide's notes"]}
           every slide also takes "notes" and "credits" (both go to the notes page).
           a figure is an svg-library component name or a path to a .svg
           (token-driven or already baked; it is baked again for the deck theme).
  --theme  overrides spec "theme" (default light); picks the house base and
           the svg-export palette.
  --base   a corporate .pptx to build on instead of templates/deck/house-<theme>.pptx.
           Its existing slides are dropped; layouts are matched by house name,
           then by layout type (title / secHead / titleOnly / obj).

Deps (optional, deck builder only): python-pptx, lxml, cairosvg; node for bin/svg-export.js.
"""
import argparse, copy, datetime, io, json, os, re, subprocess, sys, tempfile, uuid
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK_DIR = os.path.join(ROOT, "templates", "deck")
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
DECOR_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
DECOR_EXT = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"

# layout kind -> (house layout name, OOXML layout type used for --base fallback)
LAYOUTS = {
    "title": ("Title", "title"),
    "section": ("Section", "secHead"),
    "figure": ("Figure", "titleOnly"),
    "two-figures": ("Two Figures", "titleOnly"),
    "bullets": ("Bullets", "obj"),
    # v0.19.0
    "agenda": ("Agenda", "obj"),
    "quote": ("Quote", "obj"),
    "table": ("Table", "titleOnly"),
    "kpi": ("KPI", "obj"),
    "image-text": ("Image + Text", "titleOnly"),
    "closing": ("Closing", "title"),
}
N_FIGS = {"figure": 1, "two-figures": 2, "image-text": 1}
DEFAULT_TITLES = {"agenda": "Agenda", "closing": "Thank you"}
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # PowerPoint's built-in "No Style, No Grid"
NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

PIC_XML = (
    '<p:pic xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
    '<p:nvPicPr><p:cNvPr id="{id}" name="{name}" descr="{descr}"/>'
    '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
    '<p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
    '<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
)


def bake(ref, theme, workdir):
    """svg-library name or .svg path -> renderer-safe SVG bytes (bin/svg-export.js)."""
    out = subprocess.run(
        ["node", os.path.join(ROOT, "bin", "svg-export.js"), ref, "--theme", theme, "--out", workdir],
        capture_output=True, text=True)
    if out.returncode:
        sys.exit(f"svg-export failed for {ref!r}: {out.stderr.strip() or out.stdout.strip()}")
    with open(out.stdout.strip().splitlines()[-1], "rb") as f:
        return f.read()


def svg_size(svg):
    m = re.search(rb'<svg[^>]*\bwidth="([\d.]+)"[^>]*\bheight="([\d.]+)"', svg)
    return float(m.group(1)), float(m.group(2))


def svg_label(svg):
    m = re.search(rb'aria-label="([^"]*)"', svg) or re.search(rb"<title>([^<]*)</title>", svg)
    return m.group(1).decode() if m else ""


def raster(svg, width_px):
    import cairosvg
    w, h = svg_size(svg)
    return cairosvg.svg2png(bytestring=svg, output_width=width_px, output_height=round(width_px * h / w))


def add_svg_picture(part, sp_tree, svg, png, box, name, descr="", decorative=False, back=False):
    """Add a picture (PNG blip + svgBlip extension) to any slide/layout/master part.

    box = (x, y, cx, cy) in EMU. back=True inserts behind every other shape.
    """
    _, rid_png = part.get_or_add_image_part(io.BytesIO(png))
    svg_part = Part(part.package.next_partname("/ppt/media/image%d.svg"), "image/svg+xml", part.package, svg)
    rid_svg = part.relate_to(svg_part, RT.IMAGE)
    sid = 1 + max(int(e.get("id")) for e in sp_tree.iter(qn("p:cNvPr")))
    x, y, cx, cy = (int(v) for v in box)
    pic = etree.fromstring(PIC_XML.format(id=sid, name=name, descr="", rid=rid_png, x=x, y=y, cx=cx, cy=cy))
    pic.find(".//" + qn("p:cNvPr")).set("descr", descr)  # one XML set; lxml escapes it
    if decorative:
        ext = etree.SubElement(etree.SubElement(pic.find(".//" + qn("p:cNvPr")), qn("a:extLst")), qn("a:ext"), uri=DECOR_EXT)
        etree.SubElement(ext, f"{{{DECOR_NS}}}decorative", nsmap={"adec": DECOR_NS}, val="1")
    ext = etree.SubElement(etree.SubElement(pic.find(".//" + qn("a:blip")), qn("a:extLst")), qn("a:ext"), uri=SVG_EXT)
    etree.SubElement(ext, f"{{{SVG_NS}}}svgBlip", nsmap={"asvg": SVG_NS}).set(qn("r:embed"), rid_svg)
    if back:
        sp_tree.insert(2, pic)  # after nvGrpSpPr + grpSpPr
    else:
        sp_tree.append(pic)
    return pic


def contain(box, aspect):
    x, y, cx, cy = box
    w, h = (cx, cx / aspect) if cx / cy < aspect else (cy * aspect, cy)
    return x + (cx - w) / 2, y + (cy - h) / 2, w, h


def pick_layout(prs, kind):
    name, typ = LAYOUTS[kind]
    layouts = list(prs.slide_layouts)
    for test in (lambda l: l.name == name, lambda l: l._element.get("type") == typ,
                 lambda l: any(p.placeholder_format.type in TITLE_TYPES for p in l.placeholders)):
        hit = next((l for l in layouts if test(l)), None)
        if hit:
            return hit
    sys.exit(f"--base has no usable layout for {kind!r}")


def drop_all_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def fill_text(tf, items, level=0):
    """items: str or nested list of str (nested list = one level deeper)."""
    first = not any(p.text for p in tf.paragraphs)
    for it in items:
        if isinstance(it, list):
            fill_text(tf, it, level + 1)
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text, p.level = str(it), level


def norm_figures(s):
    figs = s.get("figures") or ([s["figure"]] if "figure" in s else [])
    figs = [{"svg": f} if isinstance(f, str) else dict(f) for f in figs]
    if len(figs) == 1 and "caption" in s and "caption" not in figs[0]:
        figs[0]["caption"] = s["caption"]
    if len(figs) == 1 and "alt" in s and "alt" not in figs[0]:
        figs[0]["alt"] = s["alt"]
    return figs


def lines(v):
    return [v] if isinstance(v, str) else list(v or [])


def set_lines(ph, items):
    """One paragraph per string, no level change (contact lines, agenda items)."""
    tf = ph.text_frame
    for i, t in enumerate(items):
        (tf.paragraphs[0] if i == 0 else tf.add_paragraph()).text = str(t)


def no_bullet(tf):
    for p in tf.paragraphs:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        if ppr.find(qn("a:buNone")) is None:
            etree.SubElement(ppr, qn("a:buNone"))  # fresh paragraphs: pPr has no other children


def area_below(title, sw, sh, cols, reserve=0):
    """--base fallback: split the free area under the title into `cols` columns."""
    top = (title.top + title.height if title is not None else sh // 10) + sh // 30
    gap, margin = sw // 40, sw // 20
    w = (sw - 2 * margin - gap * (cols - 1)) // cols
    return [(margin + i * (w + gap), top, w, sh - top - sh // 20 - reserve) for i in range(cols)]


def place_figures(slide, figs, boxes, caps, theme, work, sw, sh):
    for i, f in enumerate(figs):
        svg = bake(f["svg"], theme, work)
        aw, ah = svg_size(svg)
        x, y, cx, cy = contain(boxes[i], aw / ah)
        png = raster(svg, 2400 if cx > sw / 2 else 1800)
        add_svg_picture(slide.part, slide.shapes._spTree, svg, png, (x, y, cx, cy),
                        name=f"Figure {i + 1}", descr=f.get("alt") or svg_label(svg))
        if f.get("caption"):
            if i < len(caps):
                caps[i].text = f["caption"]
            else:
                tb = slide.shapes.add_textbox(Emu(boxes[i][0]), Emu(int(y + cy)), Emu(boxes[i][2]), Emu(sh // 14))
                tb.text_frame.text = f["caption"]
                tb.text_frame.word_wrap = True


def fill_table(slide, box, header, rows):
    x, y, w, h = box
    n = len(rows) + 1
    row_h = min(h // n, 54 * 9525)
    gf = slide.shapes.add_table(n, len(header), Emu(x), Emu(y), Emu(w), Emu(row_h * n))
    tbl = gf.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_TABLE  # fills below come from theme slots
    for r, values in enumerate([header] + rows):
        for c, v in enumerate(values):
            cell = tbl.cell(r, c)
            cell.text = str(v)
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(16 if r == 0 else 14)
            run.font.bold = r == 0
            run.font.color.theme_color = MSO_THEME_COLOR.LIGHT_1 if r == 0 else MSO_THEME_COLOR.TEXT_1
            if r == 0 or r % 2 == 0:  # header strip + zebra (every second body row)
                cell.fill.solid()
                cell.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_6 if r == 0 else MSO_THEME_COLOR.BACKGROUND_2
            else:
                cell.fill.background()
    return gf


def delta_slot(text):
    t = text.strip()
    return "accent4" if t[:1] in ("-", "−", "▼") else "accent5" if t[:1] in ("+", "▲") else "accent3"


def set_fill(ph, scheme):
    sppr = ph._element.spPr
    for old in sppr.findall(qn("a:solidFill")):
        sppr.remove(old)
    fill = etree.SubElement(sppr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr"), val=scheme)
    # spPr child order: xfrm, geometry, fill ... -> move the fill right after geometry/xfrm
    anchor = sppr.find(qn("a:prstGeom")) if sppr.find(qn("a:prstGeom")) is not None else sppr.find(qn("a:xfrm"))
    if anchor is not None:
        anchor.addnext(fill)


def add_fields(slide, n, spec):
    """Copy the layout's date / footer / slide-number placeholders onto the slide
    and fill them with native fields (slide number always live; date live if 'auto')."""
    want = {PP_PLACEHOLDER.SLIDE_NUMBER: spec.get("slide_numbers", True),
            PP_PLACEHOLDER.FOOTER: spec.get("footer"), PP_PLACEHOLDER.DATE: spec.get("date")}
    tree = slide.shapes._spTree
    for lph in slide.slide_layout.placeholders:
        kind = lph.placeholder_format.type
        if kind not in want or not want[kind]:
            continue
        sid = 1 + max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr")))
        nv = copy.deepcopy(lph._element.find(qn("p:nvSpPr")))
        nv.find(qn("p:cNvPr")).set("id", str(sid))
        if kind == PP_PLACEHOLDER.SLIDE_NUMBER:
            run = f'<a:fld id="{{{uuid.uuid4()}}}" type="slidenum"><a:rPr lang="en-US"/><a:t>{n}</a:t></a:fld>'
        elif kind == PP_PLACEHOLDER.DATE and want[kind] == "auto":
            run = (f'<a:fld id="{{{uuid.uuid4()}}}" type="datetime1"><a:rPr lang="en-US"/>'
                   f'<a:t>{datetime.date.today():%m/%d/%Y}</a:t></a:fld>')
        else:
            run = f'<a:r><a:rPr lang="en-US"/><a:t>{escape(str(want[kind]))}</a:t></a:r>'
        sp = etree.fromstring(
            f'<p:sp {NS}><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p>{run}</a:p></p:txBody></p:sp>')
        sp.insert(0, nv)
        tree.append(sp)


def build(spec, out, theme=None, base=None):
    theme = theme or spec.get("theme", "light")
    if theme not in ("light", "dark"):
        sys.exit(f"theme must be light or dark, not {theme!r}")
    prs = Presentation(base or os.path.join(DECK_DIR, f"house-{theme}.pptx"))
    drop_all_slides(prs)
    if spec.get("title"):
        prs.core_properties.title = spec["title"]
    sw, sh = prs.slide_width, prs.slide_height
    work = tempfile.mkdtemp(prefix="pptxdeck-")
    slides = spec["slides"]

    for n, s in enumerate(slides, 1):
        kind = s.get("layout")

        def need(ok, msg):
            if not ok:
                sys.exit(f"slide {n} ({kind}): {msg}")

        need(kind in LAYOUTS, f"layout must be one of {sorted(LAYOUTS)}")
        figs = norm_figures(s)
        need(len(figs) == N_FIGS.get(kind, 0), f"expected {N_FIGS.get(kind, 0)} figure(s), got {len(figs)}")
        slide = prs.slides.add_slide(pick_layout(prs, kind))
        phs = list(slide.placeholders)
        by_idx = {p.placeholder_format.idx: p for p in phs}
        title = next((p for p in phs if p.placeholder_format.type in TITLE_TYPES), None)
        title_text = s.get("title") or DEFAULT_TITLES.get(kind)
        if title_text:
            need(title is not None, f"layout {slide.slide_layout.name!r} has no title placeholder")
            title.text = title_text
        body = [p for p in phs if p is not title and p.placeholder_format.type in
                (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE)]
        pics = sorted((p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE), key=lambda p: p.left)

        if kind in ("title", "section") and s.get("subtitle") and body:
            body[0].text = s["subtitle"]
        elif kind == "bullets" and body:
            fill_text(body[0].text_frame, s.get("bullets", []))
        elif kind in ("figure", "two-figures"):
            # caption placeholders = body placeholders, left to right (house layouts: idx 14/16)
            caps = sorted(body, key=lambda p: p.left)
            boxes = ([(p.left, p.top, p.width, p.height) for p in pics] if len(pics) >= len(figs)
                     else area_below(title, sw, sh, len(figs), sh // 12 if len(caps) < len(figs) else 0))
            place_figures(slide, figs, boxes, caps, theme, work, sw, sh)
        elif kind == "agenda":
            items = lines(s.get("items"))
            need(1 <= len(items) <= 6, f"items: 1-6 agenda rows (designed for 3-6), got {len(items)}")
            set_lines(body[0], items)
        elif kind == "quote":
            need(bool(s.get("quote")), "quote text is required")
            set_lines(body[0], [s["quote"]])
            if s.get("attribution"):
                who = "— " + s["attribution"]
                if 13 in by_idx:
                    by_idx[13].text = who
                else:
                    body[0].text_frame.add_paragraph().text = who
        elif kind == "table":
            header, rows = s.get("header"), s.get("rows")
            need(isinstance(header, list) and 1 <= len(header) <= 8, "header: list of 1-8 column names")
            need(isinstance(rows, list) and 1 <= len(rows) <= 12, "rows: list of 1-12 rows")
            need(all(isinstance(r, list) and len(r) == len(header) for r in rows),
                 f"every row needs exactly {len(header)} cells")
            tph = by_idx.get(13)
            box = (tph.left, tph.top, tph.width, tph.height) if tph is not None else area_below(title, sw, sh, 1)[0]
            if tph is not None:
                remove(tph)
            fill_table(slide, box, header, rows)
        elif kind == "kpi":
            stats = s.get("stats", [])
            need(bool(s.get("value")), "value (the hero number) is required")
            need(len(stats) <= 3 and all(isinstance(t, dict) and t.get("value") for t in stats),
                 "stats: up to 3 {value, label} objects")
            if 13 in by_idx:  # house KPI layout
                by_idx[13].text = str(s["value"])
                if s.get("delta"):
                    by_idx[14].text = str(s["delta"])
                    set_fill(by_idx[14], delta_slot(str(s["delta"])))
                if s.get("label"):
                    by_idx[15].text = s["label"]
                for k, t in enumerate(stats):
                    by_idx[16 + 2 * k].text = str(t["value"])
                    if t.get("label"):
                        by_idx[17 + 2 * k].text = t["label"]
            else:  # --base: hero + stats as plain lines in the body
                hero = " — ".join(str(v) for v in (s["value"], s.get("label")) if v)
                set_lines(body[0], [hero + (f" ({s['delta']})" if s.get("delta") else "")] +
                          [" — ".join(str(v) for v in (t["value"], t.get("label")) if v) for t in stats])
        elif kind == "image-text":
            side = s.get("side", "left")
            need(side in ("left", "right"), "side must be left or right")
            text = s.get("bullets") if s.get("bullets") is not None else s.get("text")
            need(bool(text), "text (string) or bullets (list) is required")
            tph = by_idx.get(14)
            if pics and tph is not None:
                pbox = (pics[0].left, pics[0].top, pics[0].width, pics[0].height)
                tbox = (tph.left, tph.top, tph.width, tph.height)
            else:
                pbox, tbox = area_below(title, sw, sh, 2)
                tph = slide.shapes.add_textbox(*(Emu(v) for v in tbox))
                tph.text_frame.word_wrap = True
            if side == "right":  # mirror both boxes around the slide centre
                pbox = (sw - pbox[0] - pbox[2],) + pbox[1:]
                tbox = (sw - tbox[0] - tbox[2],) + tbox[1:]
                tph.left, tph.top, tph.width, tph.height = (Emu(v) for v in tbox)
            if isinstance(text, str):
                set_lines(tph, [text])
                no_bullet(tph.text_frame)
            else:
                fill_text(tph.text_frame, text)
            place_figures(slide, figs, [pbox], [], theme, work, sw, sh)
            for p in pics:
                remove(p)
        elif kind == "closing" and body:
            contact = lines(s.get("contact") or s.get("subtitle"))
            if contact:
                set_lines(body[0], contact)
        for p in pics:
            if p._element.getparent() is not None:
                remove(p)

        # placeholders left empty would show "Click to add ..." in edit view
        for p in slide.placeholders:
            if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE or (p.has_text_frame and not p.text_frame.text.strip()):
                remove(p)
        if kind not in ("title", "section", "closing"):
            add_fields(slide, n, spec)
        notes = [s["notes"]] if s.get("notes") else []
        credits = lines(s.get("credits")) + (lines(spec.get("credits")) if n == len(slides) else [])
        if credits:
            notes.append("Credits:\n" + "\n".join(f"- {c}" for c in credits))
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)

    prs.save(out)
    return out


def load_spec(path):
    with open(path, encoding="utf-8") as f:
        if path.endswith((".yaml", ".yml")):
            import yaml
            return yaml.safe_load(f)
        return json.load(f)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("spec")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--theme", choices=("light", "dark"))
    ap.add_argument("--base", help="corporate .pptx to build on (default: house base for the theme)")
    a = ap.parse_args()
    print(build(load_spec(a.spec), a.out, a.theme, a.base))
