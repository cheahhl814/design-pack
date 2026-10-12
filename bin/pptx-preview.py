#!/usr/bin/env python3
"""pptx-preview.py: proxy preview of a .pptx in headless Chrome, no Office needed.

Rebuilds each slide as HTML at 96 dpi (EMU / 9525 = px): master and layout decor,
slide pictures (the embedded SVG if present, else the PNG), and text boxes with
size/colour/weight/alignment/bullets resolved through slide -> layout -> master
placeholder lstStyles and the master txStyles. Then it screenshots every slide.
Text that overflows its box gets a red dashed outline.

This is a PROXY, not PowerPoint's renderer: autofit, kerning, line breaking and
font substitution differ. It catches composition bugs (overlaps, empty or
misplaced boxes, overflow, wrong theme colours); the real-renderer check is a
manual once-per-theme-release gate (see templates/deck/README.md).

    python3 bin/pptx-preview.py deck.pptx --out DIR    # DIR/slide-01.html + .png ...
"""
import argparse, base64, html, os, re, shutil, subprocess, sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

PX = 9525
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
# date / footer / slide-number placeholders take the master otherStyle, not bodyStyle
FIELD_TYPES = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)


class Theme:
    def __init__(self, prs):
        master = prs.slide_master
        root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        self.colors = {}
        for el in root.find(f".//{A}clrScheme"):
            c = el[0]
            self.colors[etree.QName(el).localname] = c.get("val") if c.tag == A + "srgbClr" else c.get("lastClr")
        self.fonts = {k: root.find(f".//{A}{tag}/{A}latin").get("typeface")
                      for k, tag in (("mj", "majorFont"), ("mn", "minorFont"))}
        self.clrmap = dict(master._element.find(qn("p:clrMap")).attrib)
        self.styles = master._element.find(qn("p:txStyles"))

    def color(self, el):
        """a:solidFill (or anything holding a colour child) -> '#rrggbb'."""
        if el is None:
            return None
        for c in el.iter(qn("a:srgbClr"), qn("a:schemeClr")):
            v = c.get("val")
            if c.tag == qn("a:schemeClr"):
                v = self.colors.get(self.clrmap.get(v, v), "000000")
            # values come from the (untrusted) pptx and land in style attrs / <style>
            return "#" + v if v and re.fullmatch(r"[0-9A-Fa-f]{6}", v) else None
        return None

    def font(self, typeface):
        if typeface and typeface.startswith("+"):
            typeface = self.fonts.get(typeface[1:3])
        # whitelist: the name is interpolated into a style attribute
        return typeface if typeface and re.fullmatch(r"[\w .-]{1,64}", typeface) else None


def lvl_chain(shape, theme, lvl):
    """lvlNpPr elements, highest priority first."""
    lists = [shape._element.find(".//" + qn("a:lstStyle"))]
    style = "p:otherStyle"
    if shape.is_placeholder:
        kind = shape.placeholder_format.type
        style = ("p:titleStyle" if kind in TITLES else "p:otherStyle" if kind in FIELD_TYPES else "p:bodyStyle")
        base = shape
        while True:
            base = getattr(base, "_base_placeholder", None)
            if base is None:
                break
            lists.append(base._element.find(".//" + qn("a:lstStyle")))
    lists.append(theme.styles.find(qn(style)) if theme.styles is not None else None)
    return [l.find(qn(f"a:lvl{lvl}pPr")) for l in lists if l is not None and l.find(qn(f"a:lvl{lvl}pPr")) is not None]


def first(els, path, attr=None):
    for e in els:
        if e is None:
            continue
        hit = e.find(path) if path else e
        if hit is not None and (attr is None or hit.get(attr) is not None):
            return hit if attr is None else hit.get(attr)
    return None


def anchor_of(shape):
    base = shape
    while base is not None:
        bp = base._element.find(".//" + qn("a:bodyPr"))
        if bp is not None and bp.get("anchor"):
            return bp.get("anchor")
        base = getattr(base, "_base_placeholder", None)
    return "t"


AUTONUM = {"arabicPeriod": "{}.", "arabicParenR": "{})", "arabicPlain": "{}"}


def text_html(shape, theme, scale):
    paras = []
    counters = {}
    for p in shape.text_frame.paragraphs:
        if not p.text:
            continue
        lvl = p.level + 1
        ppr = p._p.find(qn("a:pPr"))
        chain = [ppr] + lvl_chain(shape, theme, lvl)
        defs = [c.find(qn("a:defRPr")) if c is not None else None for c in chain]
        rpr = [r._r.find(qn("a:rPr")) for r in p.runs[:1]] + defs
        sz = int(first(rpr, None, "sz") or 1800) / 100
        bold = first(rpr, None, "b") in ("1", "true")
        italic = first(rpr, None, "i") in ("1", "true")
        color = theme.color(first(rpr, qn("a:solidFill"))) or theme.color_default
        face = theme.font(first(rpr, qn("a:latin"), "typeface")) or theme.font("+mn") or "Arial"
        algn = {"ctr": "center", "r": "right", "just": "justify"}.get(first(chain, None, "algn"), "left")
        marl = int(first(chain, None, "marL") or 0) / PX * scale
        indent = int(first(chain, None, "indent") or 0) / PX * scale
        bullet = ""
        for c in chain:
            if c is None:
                continue
            if c.find(qn("a:buNone")) is not None:
                break
            ch, num = c.find(qn("a:buChar")), c.find(qn("a:buAutoNum"))
            if ch is not None or num is not None:
                bc = theme.color(first(chain, qn("a:buClr"))) or color
                if num is not None:
                    counters[lvl] = counters.get(lvl, int(num.get("startAt", 1)) - 1) + 1
                    glyph = AUTONUM.get(num.get("type"), "{}.").format(counters[lvl])
                else:
                    glyph = ch.get("char")
                bullet = (f'<span style="display:inline-block;width:{-indent}px;text-indent:0;color:{bc};'
                          f'font-weight:700">{html.escape(glyph)}</span>')
                break
        paras.append(
            f'<p style="margin:0 0 {sz * 0.35 * scale:.1f}px 0;padding-left:{marl}px;text-indent:{indent}px;'
            f'font:{"italic " if italic else ""}{"700" if bold else "400"} {sz * 96 / 72 * scale:.1f}px/1.15 \'{face}\',sans-serif;'
            f'color:{color};text-align:{algn}">{bullet}{html.escape(p.text)}</p>')
    return "".join(paras)


def media_uri(part, pic_el):
    blip = pic_el.find(".//" + qn("a:blip"))
    svg = blip.find(f".//{{{SVG_NS}}}svgBlip")
    rid, mime = (svg.get(qn("r:embed")), "image/svg+xml") if svg is not None else (blip.get(qn("r:embed")), "image/png")
    return f"data:{mime};base64," + base64.b64encode(part.related_part(rid).blob).decode()


def box(shape, scale):
    return (f"left:{shape.left / PX * scale:.1f}px;top:{shape.top / PX * scale:.1f}px;"
            f"width:{shape.width / PX * scale:.1f}px;height:{shape.height / PX * scale:.1f}px")


def inherited(shape, path):
    """First match of `path` under the shape's own element, then its base placeholders."""
    base = shape
    while base is not None:
        hit = base._element.find(path)
        if hit is not None:
            return hit
        base = getattr(base, "_base_placeholder", None)
    return None


def shape_style(shape, theme, scale):
    """Fill + rounded geometry + bodyPr insets (placeholders inherit them from the layout)."""
    css = []
    fill = inherited(shape, f"{qn('p:spPr')}/{qn('a:solidFill')}")
    if theme.color(fill):
        css.append(f"background:{theme.color(fill)}")
    geom = inherited(shape, f"{qn('p:spPr')}/{qn('a:prstGeom')}")
    if geom is not None and geom.get("prst") == "roundRect":
        css.append("border-radius:999px")
    pad = []
    for attr, default in (("tIns", 45720), ("rIns", 91440), ("bIns", 45720), ("lIns", 91440)):
        bp = inherited(shape, f".//{qn('a:bodyPr')}[@{attr}]")
        pad.append(f"{int(bp.get(attr) if bp is not None else default) / PX * scale:.1f}px")
    css.append("padding:" + " ".join(pad))
    return ";".join(css)


def table_html(gf, theme, scale):
    rows = []
    for r, row in enumerate(gf.table.rows):
        cells = []
        for cell in row.cells:
            tcpr = cell._tc.find(qn("a:tcPr"))
            bg = theme.color(tcpr.find(qn("a:solidFill"))) if tcpr is not None else None
            run = next((rn for p in cell.text_frame.paragraphs for rn in p.runs), None)
            rpr = run._r.find(qn("a:rPr")) if run is not None else None
            sz = int(rpr.get("sz") or 1800) / 100 if rpr is not None else 18
            bold = rpr is not None and rpr.get("b") in ("1", "true")
            color = (theme.color(rpr.find(qn("a:solidFill"))) if rpr is not None else None) or theme.color_default
            cells.append(f'<td style="background:{bg or "transparent"};color:{color};font:{"700" if bold else "400"} '
                         f'{sz * 96 / 72 * scale:.1f}px/1.2 \'{theme.font("+mn") or "Arial"}\',sans-serif;'
                         f'padding:{4.8 * scale:.1f}px {9.6 * scale:.1f}px;height:{row.height / PX * scale:.1f}px;'
                         f'box-sizing:border-box">{html.escape(cell.text)}</td>')
        rows.append("<tr>" + "".join(cells) + "</tr>")
    return (f'<table class="tx" style="position:absolute;{box(gf, scale)};height:auto;border-collapse:collapse;'
            f'table-layout:fixed">{"".join(rows)}</table>')


def render_shapes(shapes, part, theme, scale, skip_placeholders):
    out = []
    for sh in shapes:
        if sh.shape_type == 13 or sh._element.tag == qn("p:pic"):  # picture
            out.append(f'<img style="position:absolute;{box(sh, scale)}" src="{media_uri(part, sh._element)}">')
        elif skip_placeholders and sh.is_placeholder:
            continue
        elif getattr(sh, "has_table", False) and sh.has_table:
            out.append(table_html(sh, theme, scale))
        elif sh.has_text_frame and sh.text_frame.text.strip():
            just = {"t": "flex-start", "ctr": "center", "b": "flex-end"}.get(anchor_of(sh), "flex-start")
            out.append(f'<div class="tx" style="position:absolute;{box(sh, scale)};display:flex;flex-direction:column;'
                       f'justify-content:{just};{shape_style(sh, theme, scale)}">'
                       f'<div>{text_html(sh, theme, scale)}</div></div>')
    return out


def render(path, out_dir, width_px=1280, screenshot=True):
    prs = Presentation(path)
    theme = Theme(prs)
    scale = width_px / (prs.slide_width / PX)
    W, H = round(prs.slide_width / PX * scale), round(prs.slide_height / PX * scale)
    bg_el = prs.slide_master._element.find(".//" + qn("p:bg"))
    bg = theme.color(bg_el) or "#ffffff"
    theme.color_default = theme.color(etree.fromstring(
        f'<a:schemeClr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="tx1"/>')) or "#000000"
    os.makedirs(out_dir, exist_ok=True)
    chrome = shutil.which("google-chrome") or shutil.which("chromium") or shutil.which("chromium-browser")
    files = []
    for n, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout
        layers = []
        if layout._element.get("showMasterSp") != "0":
            layers += render_shapes(layout.slide_master.shapes, layout.slide_master.part, theme, scale, True)
        layers += render_shapes(layout.shapes, layout.part, theme, scale, True)
        layers += render_shapes(slide.shapes, slide.part, theme, scale, False)
        doc = (f'<!doctype html><meta charset="utf-8"><title>slide {n}</title>'
               f'<style>html,body{{margin:0;overflow:hidden}}.slide{{position:relative;width:{W}px;height:{H}px;'
               f'background:{bg};overflow:hidden}}.tx{{box-sizing:border-box;overflow:hidden}}'
               f'.tx p:last-child{{margin-bottom:0!important}}.overflow{{outline:3px dashed red}}</style>'
               f'<div class="slide">{"".join(layers)}</div>'
               '<script>document.querySelectorAll(".tx").forEach(e=>{if(e.scrollHeight>e.clientHeight+1)'
               'e.classList.add("overflow")})</script>')
        html_path = os.path.join(out_dir, f"slide-{n:02d}.html")
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(doc)
        files.append(html_path)
        if screenshot:
            if not chrome:
                sys.exit("no google-chrome/chromium on PATH; rerun with --no-screenshot")
            subprocess.run([chrome, "--headless", "--disable-gpu", "--hide-scrollbars",
                            f"--window-size={W},{H}", f"--screenshot={html_path[:-5]}.png",
                            "file://" + os.path.abspath(html_path)], capture_output=True, check=True)
    return files


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("pptx")
    ap.add_argument("--out", required=True)
    ap.add_argument("--width", type=int, default=1280, help="preview width in px (default 1280)")
    ap.add_argument("--no-screenshot", action="store_true", help="write HTML only")
    a = ap.parse_args()
    for f in render(a.pptx, a.out, a.width, not a.no_screenshot):
        print(f)
